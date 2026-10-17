import csv
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .catalog import Catalog, ImportReport


REQUIRED = ("sku", "name", "shade", "price_uah_per_m", "stock_m", "updated_at")
PRIVATE = ("purchase_price_uah_per_m", "supplier")


class ImportError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


def _rows(path: Path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            yield from csv.reader(stream, strict=True)
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            yield from workbook.worksheets[0].iter_rows(values_only=True)
        finally:
            workbook.close()
    else:
        raise ImportError(["Підтримуються лише CSV та XLSX."])


def _text(value, field: str, *, optional: bool = False) -> str | None:
    if value is None or value == "":
        if optional:
            return None
        raise ValueError(f"{field}: обов'язкове поле порожнє")
    if not isinstance(value, str):
        raise ValueError(f"{field}: потрібен текст (SKU в Excel зберігайте як текст)")
    value = " ".join(value.split())
    if not value or len(value) > 200:
        raise ValueError(f"{field}: довжина має бути від 1 до 200 символів")
    return value


def _decimal(value, field: str, *, optional: bool = False) -> str | None:
    if value is None or value == "":
        if optional:
            return None
        raise ValueError(f"{field}: обов'язкове поле порожнє")
    text = str(value).strip()
    if not re.fullmatch(r"[0-9]{1,12}(?:\.[0-9]{1,2})?", text):
        raise ValueError(f"{field}: потрібне невід'ємне число, до 12 цифр перед крапкою та двох після")
    try:
        return format(Decimal(text), "f")
    except InvalidOperation:
        raise ValueError(f"{field}: некоректне число") from None


def import_catalog(catalog: Catalog, path: str | Path) -> ImportReport:
    errors = []
    products = []
    seen = set()
    rows = _rows(Path(path))
    try:
        header = next(rows, ())
        if (not header or len(set(header)) != len(header)
                or not set(REQUIRED).issubset(header)
                or set(header) - set(REQUIRED + PRIVATE)):
            raise ImportError(["Рядок 1: некоректні заголовки каталогу."])
        for number, cells in enumerate(rows, 2):
            if all(cell is None or cell == "" for cell in cells):
                continue
            try:
                if len(cells) != len(header):
                    raise ValueError("кількість полів не відповідає заголовку")
                row = dict(zip(header, cells))
                if any(isinstance(v, str) and v.lstrip().startswith("=") for v in cells):
                    raise ValueError("формули не підтримуються")
                product = {field: _text(row[field], field) for field in ("sku", "name", "shade")}
                if product["sku"] in seen:
                    raise ValueError("повтор артикула")
                seen.add(product["sku"])
                product["price_uah_per_m"] = _decimal(row["price_uah_per_m"], "price_uah_per_m")
                product["stock_m"] = _decimal(row["stock_m"], "stock_m", optional=True)
                raw_date = _text(row["updated_at"], "updated_at")
                try:
                    updated_at = datetime.fromisoformat(raw_date)
                    if updated_at.tzinfo is None or "T" not in raw_date:
                        raise ValueError
                    product["updated_at"] = updated_at.astimezone(timezone.utc).isoformat()
                except ValueError:
                    raise ValueError("updated_at: потрібна дата ISO 8601 з часовим поясом") from None
                product[PRIVATE[0]] = _decimal(row.get(PRIVATE[0]), PRIVATE[0], optional=True)
                product[PRIVATE[1]] = _text(row.get(PRIVATE[1]), PRIVATE[1], optional=True)
                products.append(product)
            except ValueError as error:
                errors.append(f"Рядок {number}: {error}.")
        # A missing export must not accidentally wipe the live catalog.
        if not products and not errors:
            errors.append("Каталог порожній; імпорт скасовано.")
        if errors:
            raise ImportError(errors)
        return catalog.replace(products)
    finally:
        rows.close()
